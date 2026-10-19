import json

from decimal import Decimal
from io import BytesIO

from django.core.files.base import ContentFile
from django.core.management import call_command
from django.test import TestCase
from openpyxl import Workbook

from apps.catalog.models import Brand, Category, Product, ProductText
from apps.corecfg.models import Language
from apps.ingestion.models import ImportJob, Offer, OfferText
from apps.ingestion.pipeline import run_import
from apps.mapping.models import CategoryMapping
from apps.pricing.models import PriceAction, PriceRule
from apps.suppliers.models import FeedRole, SourceKind, Supplier, SupplierFeed


class ImportPipelineTests(TestCase):
    def test_json_feed_creates_published_offer(self):
        supplier = Supplier.objects.create(code="demo", name="Demo")
        category = Category.objects.create(name="Meškerės")
        CategoryMapping.objects.create(
            supplier=supplier,
            source_category_raw="Rods",
            source_sub_category_raw="",
            source_path="Rods",
            category=category,
        )
        PriceRule.objects.create(
            name="20%",
            priority=10,
            action=PriceAction.MARKUP_PERCENT,
            value=Decimal("20"),
        )
        payload = json.dumps(
            [
                {
                    "id": "SKU-1",
                    "title": "Spinningas",
                    "price": "10.00",
                    "category": "Rods",
                    "brand": "Shimano",
                    "stock": 3,
                    "image": "https://cdn.example/spinningas.jpg",
                }
            ]
        ).encode()
        feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="JSON demo",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
            config={
                "fields": {
                    "external_id": "id",
                    "name": "title",
                    "purchase_price": "price",
                    "category_path": "category",
                    "brand_name": "brand",
                    "stock": "stock",
                    "images": "image",
                }
            },
        )
        feed.file.save("items.json", ContentFile(payload), save=True)
        job = ImportJob.objects.create(supplier=supplier, feed=feed)
        run_import(job)

        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.Status.SUCCESS)
        offer = Offer.objects.get(supplier=supplier, external_id="SKU-1")
        self.assertEqual(offer.name, "Spinningas")
        self.assertEqual(offer.category, category)
        self.assertEqual(offer.selling_price, Decimal("12.00"))
        self.assertTrue(offer.is_published)
        self.assertIsNotNone(offer.product_id)

    def test_fox_adapter_uses_text_and_media_and_brand_fallback(self):
        supplier = Supplier.objects.create(code="fox", name="FOX")
        category = Category.objects.create(name="Accessories")
        CategoryMapping.objects.create(
            supplier=supplier,
            source_category_raw="Accessories",
            source_sub_category_raw="Float",
            source_path="Accessories > Float",
            category=category,
        )

        workbook = Workbook()
        base_ws = workbook.active
        base_ws.title = "BASE"
        base_ws.append(
            [
                "Art-No.",
                "Product Group",
                "Brand",
                "Name",
                "Trade EUR",
                "RRP EUR",
                "Barcode Single EAN",
                "Pack Quantity",
                "Category",
                "Sub Category",
                "Family",
            ]
        )
        base_ws.append(
            [
                "AAA-1",
                "GRP-1",
                "Fox Rage",
                "Variant One",
                "10.50",
                "20.99",
                "1111111111111",
                "5",
                "Accessories",
                "Float",
                "Family A",
            ]
        )
        base_ws.append(
            [
                "AAA-2",
                "GRP-1",
                "",
                "Variant Two",
                "12.00",
                "22.49",
                "2222222222222",
                "3",
                "Accessories",
                "Float",
                "Family A",
            ]
        )

        text_ws = workbook.create_sheet("TEXT")
        text_ws.append(
            [
                "Art-No.",
                "Product Group",
                "Brand",
                "Product Description",
                "Techn. Description",
                "Combined Full Text",
            ]
        )
        text_ws.append(
            [
                "AAA-1",
                "GRP-1",
                "Fox Rage",
                "Group Name",
                "Variant One Specs",
                "Line 1\nLine 2",
            ]
        )
        text_ws.append(
            [
                "AAA-2",
                "GRP-1",
                "",
                "Group Name",
                "Variant Two Specs",
                "Combined text second",
            ]
        )

        media_ws = workbook.create_sheet("MEDIA CLOUD LINKS")
        media_ws.append(["Art-No.", "Main", "Studio1"])
        media_ws.append(["AAA-1", "https://img.example.com/aaa-1-main.jpg", ""])
        media_ws.append(
            [
                "AAA-2",
                "https://img.example.com/aaa-2-main.jpg",
                "https://img.example.com/aaa-2-studio.jpg",
            ]
        )

        workbook.create_sheet("MEDIA FILES REFERENCE").append(["Art-No.", "Main"])
        workbook["MEDIA FILES REFERENCE"].append(["AAA-1", "aaa-1-main.jpg"])
        workbook["MEDIA FILES REFERENCE"].append(["AAA-2", "aaa-2-main.jpg"])

        specs_ws = workbook.create_sheet("SPECS")
        specs_ws.append(
            [
                "Art-No.",
                "Product Group",
                "Brand",
                "Name",
                "Size",
                "Unit Height",
                "Retail Height cm",
                "Retail Width cm",
                "Retail Depth cm",
                "Retail Weight gram",
            ]
        )
        specs_ws.append(["AAA-1", "GRP-1", "Fox Rage", "Variant One", "Medium", 10, 2.5, 29.6, 19.8, 178])
        specs_ws.append(["AAA-2", "GRP-1", "Fox Rage", "Variant Two", "Large", 12, 2.6, 30.0, 20.1, 180])

        payload = BytesIO()
        workbook.save(payload)

        feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="FOX demo",
            kind=SourceKind.EXCEL_UPLOAD,
            adapter_code="fox",
        )
        feed.file.save("fox.xlsx", ContentFile(payload.getvalue()), save=True)
        job = ImportJob.objects.create(supplier=supplier, feed=feed)
        run_import(job)

        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.Status.SUCCESS)

        first = Offer.objects.get(supplier=supplier, external_id="AAA-1")
        second = Offer.objects.get(supplier=supplier, external_id="AAA-2")

        # base_name = Product Description + brand prefiksas
        self.assertEqual(first.base_name, "Fox Rage Group Name")
        self.assertEqual(first.size_raw, "Medium")
        self.assertEqual(first.shipping_height_cm, Decimal("2.5"))
        self.assertEqual(first.shipping_width_cm, Decimal("29.6"))
        self.assertEqual(first.shipping_depth_cm, Decimal("19.8"))
        self.assertEqual(first.shipping_weight_kg, Decimal("0.178"))
        self.assertEqual(first.description, "Line 1\nLine 2")
        self.assertEqual(first.combined_full_text, "Line 1\nLine 2")
        self.assertEqual(first.category_raw, "Accessories")
        self.assertEqual(first.sub_category_raw, "Float")
        self.assertEqual(first.family_raw, "Family A")
        self.assertEqual(first.product_id, second.product_id)
        self.assertEqual(second.size_raw, "Large")
        self.assertEqual(first.product.size_raw, "Medium")
        self.assertEqual(first.product.shipping_height_cm, Decimal("2.5"))
        self.assertEqual(first.product.shipping_width_cm, Decimal("29.6"))
        self.assertEqual(first.product.shipping_depth_cm, Decimal("19.8"))
        self.assertEqual(first.product.shipping_weight_kg, Decimal("0.178"))
        self.assertEqual(first.rrp_price, Decimal("20.99"))
        # Pardavimo kaina = RRP EUR (su PVM); savikaina = Trade EUR be PVM
        self.assertEqual(first.selling_price, Decimal("20.99"))
        self.assertEqual(first.category, category)
        self.assertEqual(first.extra["identifiers"]["product_group"], "GRP-1")
        self.assertEqual(first.images, ["https://img.example.com/aaa-1-main.jpg"])
        self.assertEqual(first.main_image_url, "https://img.example.com/aaa-1-main.jpg")
        # Unit Height -> unit_height axis (gaminio matmuo, gali būti variantas)
        self.assertEqual(first.attributes["axes"]["unit_height"], 10)
        self.assertNotIn("Unit Height", first.attributes.get("features", {}))
        # Retail* matmenys -> kanoniniai Offer modelio laukai (Width/Height/Depth cm, Weight kg)
        self.assertEqual(first.shipping_width_cm, Decimal("29.6"))
        self.assertEqual(first.shipping_height_cm, Decimal("2.5"))
        self.assertEqual(first.shipping_depth_cm, Decimal("19.8"))
        self.assertEqual(first.shipping_weight_kg, Decimal("0.178"))
        self.assertNotIn("shipping_dimensions", first.attributes)
        # SPECS klasifikuota: Size -> 'size', Unit Height -> 'unit_height'
        self.assertEqual(
            first.attributes["axes"], {"size": "Medium", "unit_height": 10}
        )
        self.assertNotIn("specs", first.attributes)
        self.assertNotIn("image_files", first.attributes)

        self.assertEqual(second.brand_name_raw, "Fox Rage")
        self.assertEqual(second.brand.name, "Fox Rage")
        self.assertIsNotNone(second.product)
        self.assertEqual(second.product.brand.name, "Fox Rage")
        self.assertEqual(second.product.main_image_url, "https://img.example.com/aaa-2-main.jpg")
        self.assertEqual(second.product.base_name, "Fox Rage Group Name")
        self.assertEqual(
            second.images,
            [
                "https://img.example.com/aaa-2-main.jpg",
                "https://img.example.com/aaa-2-studio.jpg",
            ],
        )
        self.assertEqual(
            second.product.images,
            [
                "https://img.example.com/aaa-2-main.jpg",
                "https://img.example.com/aaa-2-studio.jpg",
            ],
        )
        self.assertNotIn("image_files", second.product.attributes)
        self.assertTrue(Brand.objects.filter(name="Fox Rage").exists())

    def test_fox_adapter_pack_fields_and_ean_fallback(self):
        """Trade Pack laukai + barkodų fallback.

        - Abu EAN -> ean=single, pack_barcode=tradepack
        - Tik single -> sold_as_pack=False
        - Tik tradepack -> ean=pack, sold_as_pack=True ('dėžė kaip prekė')
        """
        supplier = Supplier.objects.create(code="fox", name="FOX")

        workbook = Workbook()
        base_ws = workbook.active
        base_ws.title = "BASE"
        base_ws.append(
            [
                "Art-No.",
                "Product Group",
                "Brand",
                "Name",
                "Trade EUR",
                "RRP EUR",
                "Barcode Single EAN",
                "Barcode Tradepack EAN",
                "Pack Unit",
                "Pack Quantity",
                "Trade Pack QTY",
                "Category",
            ]
        )
        base_ws.append(
            [
                "AAA-1", "GRP-1", "Fox Rage", "Both EAN",
                "3.5", "6.99", "1111111111111", "5555555555551",
                "pcs", 1, 5, "Accessories",
            ]
        )
        base_ws.append(
            [
                "AAA-2", "GRP-1", "Fox Rage", "Single only",
                "4.11", "8.46", "2222222222222", "",
                "pcs", 1, 1, "Accessories",
            ]
        )
        base_ws.append(
            [
                "NZS-1", "GRP-2", "Fox Rage", "Pack only 60pcs",
                "0.5", "1.0", "", "5056212116980",
                "pcs", 1, 60, "Accessories",
            ]
        )

        payload = BytesIO()
        workbook.save(payload)
        feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="FOX pack demo",
            kind=SourceKind.EXCEL_UPLOAD,
            adapter_code="fox",
        )
        feed.file.save("fox_pack.xlsx", ContentFile(payload.getvalue()), save=True)
        job = ImportJob.objects.create(supplier=supplier, feed=feed)
        run_import(job)
        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.Status.SUCCESS)

        both = Offer.objects.get(supplier=supplier, external_id="AAA-1")
        self.assertEqual(both.ean, "1111111111111")
        self.assertEqual(both.pack_barcode, "5555555555551")
        self.assertEqual(both.pack_unit, "pcs")
        self.assertEqual(both.pack_quantity, 1)
        self.assertEqual(both.trade_pack_qty, 5)
        self.assertFalse(both.sold_as_pack)

        single = Offer.objects.get(supplier=supplier, external_id="AAA-2")
        self.assertEqual(single.ean, "2222222222222")
        self.assertEqual(single.pack_barcode, "")
        self.assertFalse(single.sold_as_pack)

        pack_only = Offer.objects.get(supplier=supplier, external_id="NZS-1")
        self.assertEqual(pack_only.ean, "5056212116980")
        self.assertEqual(pack_only.pack_barcode, "5056212116980")
        self.assertEqual(pack_only.trade_pack_qty, 60)
        self.assertTrue(pack_only.sold_as_pack)

    def test_import_feed_command_imports_all_active_feeds_and_preserves_base_data(self):
        supplier = Supplier.objects.create(code="fox", name="FOX")
        category = Category.objects.create(name="Accessories")
        CategoryMapping.objects.create(
            supplier=supplier,
            source_category_raw="Accessories",
            source_sub_category_raw="Float",
            source_path="Accessories > Float",
            category=category,
        )

        base_payload = json.dumps(
            [
                {
                    "id": "AAA-1",
                    "title": "Base Product Name",
                    "price": "10.00",
                    "category": "Accessories > Float",
                    "brand": "Fox",
                    "stock": 3,
                }
            ]
        ).encode()
        base_feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="01-base",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
            config={
                "fields": {
                    "external_id": "id",
                    "name": "title",
                    "purchase_price": "price",
                    "category_path": "category",
                    "brand_name": "brand",
                    "stock": "stock",
                }
            },
        )
        base_feed.file.save("fox-base.json", ContentFile(base_payload), save=True)

        stock_feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="02-stock",
            kind=SourceKind.CSV_UPLOAD,
            role=FeedRole.STOCK,
            adapter_code="fox_stock",
        )
        stock_feed.file.save(
            "fox-stock.csv",
            ContentFile("Art-No.;Status;Date;EAN\nAAA-1;Y;;1111111111111\n".encode("utf-8")),
            save=True,
        )

        discontinued_feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="03-discontinued",
            kind=SourceKind.CSV_UPLOAD,
            role=FeedRole.DISCONTINUED,
            adapter_code="fox_discontinued",
        )
        discontinued_feed.file.save(
            "fox-discontinued.csv",
            ContentFile("Art-No.\nAAA-1\n".encode("utf-8")),
            save=True,
        )

        call_command("import_feed", "--supplier", "fox")

        offer = Offer.objects.get(supplier=supplier, external_id="AAA-1")
        self.assertEqual(offer.name, "Base Product Name")
        self.assertEqual(offer.brand_name_raw, "Fox")
        self.assertEqual(offer.category, category)
        self.assertTrue(offer.is_discontinued)
        self.assertFalse(offer.is_available)
        self.assertEqual(offer.stock, 0)
        self.assertFalse(offer.is_published)
        self.assertEqual(
            ImportJob.objects.filter(supplier=supplier, status=ImportJob.Status.SUCCESS).count(),
            3,
        )

    def test_partial_update_feed_does_not_create_placeholder_offer_when_base_missing(self):
        supplier = Supplier.objects.create(code="fox", name="FOX")
        stock_feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="02-stock",
            kind=SourceKind.CSV_UPLOAD,
            role=FeedRole.STOCK,
            adapter_code="fox_stock",
        )
        stock_feed.file.save(
            "fox-stock.csv",
            ContentFile("Art-No.;Status;Date;EAN\nUNKNOWN-1;N;;\n".encode("utf-8")),
            save=True,
        )

        call_command("import_feed", "--supplier", "fox")

        self.assertFalse(Offer.objects.filter(supplier=supplier, external_id="UNKNOWN-1").exists())
        job = ImportJob.objects.filter(supplier=supplier).latest("id")
        self.assertEqual(job.status, ImportJob.Status.SUCCESS)
        self.assertEqual(job.stats.get("skipped_partial_missing"), 1)

    def test_stock_partial_update_matches_offer_by_ean_when_external_id_differs(self):
        supplier = Supplier.objects.create(code="fox", name="FOX")
        category = Category.objects.create(name="Accessories")
        CategoryMapping.objects.create(
            supplier=supplier,
            source_category_raw="Accessories",
            source_sub_category_raw="",
            source_path="Accessories",
            category=category,
        )

        base_payload = json.dumps(
            [
                {
                    "id": "AAA-1",
                    "title": "Base Product Name",
                    "price": "99.99",
                    "category": "Accessories",
                    "brand": "Fox",
                    "ean": "1111111111111",
                    "stock": 0,
                }
            ]
        ).encode()
        base_feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="01-base",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
            config={
                "fields": {
                    "external_id": "id",
                    "name": "title",
                    "purchase_price": "price",
                    "category_path": "category",
                    "brand_name": "brand",
                    "ean": "ean",
                    "stock": "stock",
                }
            },
        )
        base_feed.file.save("fox-base.json", ContentFile(base_payload), save=True)

        stock_feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="02-stock",
            kind=SourceKind.CSV_UPLOAD,
            role=FeedRole.STOCK,
            adapter_code="fox_stock",
        )
        stock_feed.file.save(
            "fox-stock.csv",
            ContentFile("Art-No.;Status;Date;EAN\nALT-CODE;Y;;1111111111111\n".encode("utf-8")),
            save=True,
        )

        call_command("import_feed", "--supplier", "fox")

        offer = Offer.objects.get(supplier=supplier, external_id="AAA-1")
        self.assertTrue(offer.is_available)
        self.assertEqual(offer.stock, 10)
        self.assertEqual(offer.ean, "1111111111111")
        self.assertFalse(Offer.objects.filter(supplier=supplier, external_id="ALT-CODE").exists())

    def test_moss_xml_import_links_default_and_variant_to_same_product(self):
        supplier = Supplier.objects.create(code="moss", name="MOSS")
        category = Category.objects.create(name="Fishing lines")
        CategoryMapping.objects.create(
            supplier=supplier,
            source_category_raw="Fishing lines",
            source_sub_category_raw="Braided lines - main lines > Spinning",
            source_path="Fishing lines > Braided lines - main lines > Spinning",
            category=category,
        )

        payload = """
<SHOP>
  <SHOPITEM>
    <OC>101000000</OC>
    <EAN>8586018469326</EAN>
    <VARIANT></VARIANT>
    <PRODUCT>Delphin NEXO 12 green</PRODUCT>
    <ZNACKA>DELPHIN</ZNACKA>
    <ROZMER>0,10mm 6,8kg 130m</ROZMER>
    <BALENIE>10</BALENIE>
    <HMOTNOST>5.1</HMOTNOST>
    <WEIGHT>50.1</WEIGHT>
    <WIDTH>2.5</WIDTH>
    <LENGTH>8.5</LENGTH>
    <HEIGHT>11.5</HEIGHT>
    <DESCRIPTION>Main variant description</DESCRIPTION>
    <IMG>https://delph.b-cdn.net/obr/101000000.jpg</IMG>
    <PRICE_VAT>32,95</PRICE_VAT>
    <CATEGORYTEXT>Fishing lines ||| Braided lines - main lines ||| Spinning</CATEGORYTEXT>
    <AVAILABLE>10</AVAILABLE>
  </SHOPITEM>
  <SHOPITEM>
    <OC>101000001</OC>
    <EAN>8586018469333</EAN>
    <VARIANT>101000000</VARIANT>
    <PRODUCT>Delphin NEXO 12 green</PRODUCT>
    <ZNACKA>DELPHIN</ZNACKA>
    <ROZMER>0,12mm 8,7kg 130m</ROZMER>
    <BALENIE>20</BALENIE>
    <HMOTNOST>5.5</HMOTNOST>
    <WEIGHT>52.2</WEIGHT>
    <WIDTH>2.5</WIDTH>
    <LENGTH>8.5</LENGTH>
    <HEIGHT>11.5</HEIGHT>
    <DESCRIPTION>Second variant description</DESCRIPTION>
    <IMG>https://delph.b-cdn.net/obr/101000000.jpg</IMG>
    <IMG_2>https://delph.b-cdn.net/obr/101000001_2.jpg</IMG_2>
    <PRICE_VAT>32,95</PRICE_VAT>
    <CATEGORYTEXT>Fishing lines ||| Braided lines - main lines ||| Spinning</CATEGORYTEXT>
    <AVAILABLE>0</AVAILABLE>
  </SHOPITEM>
</SHOP>
        """.strip().encode("utf-8")

        feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="MOSS feed",
            kind=SourceKind.XML_URL,
            adapter_code="moss",
        )
        feed.file.save("moss-feed.xml", ContentFile(payload), save=True)

        job = ImportJob.objects.create(supplier=supplier, feed=feed)
        run_import(job)

        job.refresh_from_db()
        self.assertEqual(job.status, ImportJob.Status.SUCCESS)
        self.assertEqual(job.stats.get("upserted"), 2)

        default_offer = Offer.objects.get(supplier=supplier, external_id="101000000")
        variant_offer = Offer.objects.get(supplier=supplier, external_id="101000001")

        self.assertEqual(default_offer.product_id, variant_offer.product_id)
        # Savikaina hipotetinė: PRICE_VAT -30%; pardavimo = PRICE_VAT (RRP, su PVM)
        self.assertEqual(default_offer.purchase_price, Decimal("23.0600"))
        self.assertEqual(default_offer.rrp_price, Decimal("32.95"))
        self.assertEqual(default_offer.selling_price, Decimal("32.95"))
        self.assertTrue(default_offer.is_available)
        self.assertEqual(default_offer.stock, 10)
        self.assertFalse(variant_offer.is_available)
        self.assertEqual(variant_offer.stock, 0)
        self.assertEqual(default_offer.category_raw, "Fishing lines")
        self.assertEqual(default_offer.sub_category_raw, "Braided lines - main lines > Spinning")
        self.assertEqual(default_offer.size_raw, "0,10mm 6,8kg 130m")
        self.assertEqual(default_offer.package_raw, "10")
        self.assertEqual(default_offer.unit_weight_raw, "5.1")
        self.assertEqual(default_offer.shipping_height_cm, Decimal("11.5"))
        self.assertEqual(default_offer.shipping_width_cm, Decimal("2.5"))
        self.assertEqual(default_offer.shipping_depth_cm, Decimal("8.5"))
        self.assertEqual(default_offer.shipping_weight_kg, Decimal("0.0501"))
        self.assertEqual(
            variant_offer.sub_category_raw,
            "Braided lines - main lines > Spinning",
        )
        self.assertEqual(variant_offer.size_raw, "0,12mm 8,7kg 130m")
        self.assertEqual(variant_offer.package_raw, "20")
        self.assertEqual(variant_offer.unit_weight_raw, "5.5")
        self.assertEqual(default_offer.name, "Delphin NEXO 12 green")
        self.assertEqual(default_offer.base_name, "Delphin NEXO 12 green")
        self.assertEqual(default_offer.extra["product_group"], "101000000")
        self.assertEqual(variant_offer.extra["product_group"], "101000000")
        self.assertEqual(default_offer.main_image_url, "https://delph.b-cdn.net/obr/101000000.jpg")
        self.assertEqual(
            variant_offer.images,
            [
                "https://delph.b-cdn.net/obr/101000000.jpg",
                "https://delph.b-cdn.net/obr/101000001_2.jpg",
            ],
        )


class LanguageFeedTests(TestCase):
    def setUp(self):
        self.lang_en = Language.objects.create(code="en", name="English", is_base=True)
        self.lang_pl = Language.objects.create(code="pl", name="Polish")
        self.supplier = Supplier.objects.create(code="fox", name="FOX")

    def _json_feed(self, name, language, rows):
        feed = SupplierFeed.objects.create(
            supplier=self.supplier,
            name=name,
            kind=SourceKind.JSON_URL,
            role=FeedRole.CATALOG,
            language=language,
            feed_group="pricelist",
            adapter_code="generic_json",
            config={
                "fields": {
                    "external_id": "id",
                    "name": "title",
                    "purchase_price": "price",
                    "description": "desc",
                    "stock": "stock",
                }
            },
        )
        feed.file.save(
            f"{name}.json",
            ContentFile(json.dumps(rows).encode()),
            save=True,
        )
        return feed

    def test_translation_feed_writes_offertext_and_producttext(self):
        self._json_feed(
            "en",
            self.lang_en,
            [{"id": "A-1", "title": "Base EN", "price": "10.00", "desc": "EN desc", "stock": 2}],
        )
        pl_feed = self._json_feed(
            "pl",
            self.lang_pl,
            [{"id": "A-1", "title": "Baza PL", "price": "10.00", "desc": "Opis PL", "stock": 2}],
        )

        call_command("import_feed", "--supplier", "fox", "--feed", str(pl_feed.pk))

        offer = Offer.objects.get(supplier=self.supplier, external_id="A-1")
        # Kalbinis feed'as neperrašo bazinių laukų
        self.assertEqual(offer.name, "Base EN")
        self.assertEqual(offer.description, "EN desc")

        text = OfferText.objects.get(offer=offer, language=self.lang_pl)
        self.assertEqual(text.name, "Baza PL")
        self.assertEqual(text.description, "Opis PL")

        ptext = ProductText.objects.get(product=offer.product, language=self.lang_pl)
        self.assertEqual(ptext.name, "Baza PL")
        self.assertTrue(ptext.slug)

    def test_translation_feed_alone_creates_minimal_offer(self):
        pl_feed = self._json_feed(
            "pl",
            self.lang_pl,
            [{"id": "A-2", "title": "Tylko PL", "desc": "Opis"}],
        )
        call_command("import_feed", "--supplier", "fox", "--feed", str(pl_feed.pk))
        offer = Offer.objects.get(supplier=self.supplier, external_id="A-2")
        self.assertFalse(offer.is_published)
        self.assertTrue(OfferText.objects.filter(offer=offer, language=self.lang_pl).exists())

    def test_group_expands_to_all_languages(self):
        en_feed = self._json_feed("en", self.lang_en, [{"id": "A-3", "title": "EN", "price": "1"}])
        self._json_feed("pl", self.lang_pl, [{"id": "A-3", "title": "PL"}])
        call_command("import_feed", "--supplier", "fox", "--feed", str(en_feed.pk))
        # Abu feed'ai importuoti (2 job'ai), bazinė kalba pirmiausia
        self.assertEqual(ImportJob.objects.filter(supplier=self.supplier).count(), 2)
        offer = Offer.objects.get(supplier=self.supplier, external_id="A-3")
        self.assertEqual(offer.name, "EN")
        self.assertTrue(OfferText.objects.filter(offer=offer, language=self.lang_pl).exists())


class ChangeDetectionTests(TestCase):
    def test_reimport_same_file_marks_unchanged(self):
        supplier = Supplier.objects.create(code="demo", name="Demo")
        payload = json.dumps([{"id": "S-1", "title": "X", "price": "5.00"}]).encode()
        feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="demo",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
            config={
                "fields": {
                    "external_id": "id",
                    "name": "title",
                    "purchase_price": "price",
                }
            },
        )
        feed.file.save("demo.json", ContentFile(payload), save=True)

        run_import(ImportJob.objects.create(supplier=supplier, feed=feed))
        job2 = run_import(ImportJob.objects.create(supplier=supplier, feed=feed))
        self.assertEqual(job2.stats.get("unchanged"), 1)

        # Pasikeitus kainai – hash kitas, pilnas upsert
        feed.file.save("demo.json", ContentFile(
            json.dumps([{"id": "S-1", "title": "X", "price": "6.00"}]).encode()
        ), save=True)
        job3 = run_import(ImportJob.objects.create(supplier=supplier, feed=feed))
        self.assertEqual(job3.stats.get("unchanged", 0), 0)
        offer = Offer.objects.get(supplier=supplier, external_id="S-1")
        self.assertEqual(offer.purchase_price, Decimal("6.00"))


class FeedDueTests(TestCase):
    def test_is_due_and_import_due_feeds(self):
        supplier = Supplier.objects.create(code="demo", name="Demo")
        feed = SupplierFeed.objects.create(
            supplier=supplier,
            name="hourly",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
            sync_interval_minutes=60,
            config={"fields": {"external_id": "id", "name": "title"}},
        )
        feed.file.save(
            "due.json", ContentFile(json.dumps([{"id": "D-1", "title": "X"}]).encode()),
            save=True,
        )
        SupplierFeed.objects.create(
            supplier=supplier,
            name="manual",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
        )
        self.assertTrue(feed.is_due())
        call_command("import_due_feeds")
        feed.refresh_from_db()
        self.assertFalse(feed.is_due())
        self.assertTrue(Offer.objects.filter(supplier=supplier, external_id="D-1").exists())


class EanMatchingTests(TestCase):
    def test_ean_matches_existing_product_across_suppliers(self):
        fox = Supplier.objects.create(code="fox", name="FOX")
        moss = Supplier.objects.create(code="moss", name="MOSS")
        product = Product.objects.create(sku="fox-1", name="X", ean="4000000000001")

        payload = json.dumps(
            [{"id": "M-1", "title": "Same product", "price": "9.99", "ean": "4000000000001"}]
        ).encode()
        feed = SupplierFeed.objects.create(
            supplier=moss,
            name="moss",
            kind=SourceKind.JSON_URL,
            adapter_code="generic_json",
            config={
                "fields": {
                    "external_id": "id",
                    "name": "title",
                    "purchase_price": "price",
                    "ean": "ean",
                }
            },
        )
        feed.file.save("moss.json", ContentFile(payload), save=True)
        run_import(ImportJob.objects.create(supplier=moss, feed=feed))

        offer = Offer.objects.get(supplier=moss, external_id="M-1")
        self.assertEqual(offer.product_id, product.id)

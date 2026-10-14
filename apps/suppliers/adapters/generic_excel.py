from __future__ import annotations

import csv
import io
from collections.abc import Iterator
from typing import Any

from openpyxl import load_workbook

from apps.suppliers.adapters.base import BaseAdapter
from apps.suppliers.adapters.registry import register_adapter


@register_adapter
class GenericExcelAdapter(BaseAdapter):
    """Excel / CSV. Config: {"sheet": "Prekes", "header_row": 1, "fields": {...}}."""

    code = "generic_excel"

    def parse(self, payload: bytes) -> Iterator[dict[str, Any]]:
        if self._looks_like_csv(payload):
            yield from self._parse_csv(payload)
            return
        yield from self._parse_xlsx(payload)

    def _looks_like_csv(self, payload: bytes) -> bool:
        name = (self.feed.file.name if self.feed.file else "").lower()
        if name.endswith(".csv"):
            return True
        if name.endswith((".xlsx", ".xlsm", ".xls")):
            return False
        return payload[:2] != b"PK"

    def _parse_xlsx(self, payload: bytes) -> Iterator[dict[str, Any]]:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
        sheet_name = self.config.get("sheet")
        worksheet = workbook[sheet_name] if sheet_name else workbook.active
        header_row = int(self.config.get("header_row", 1))
        rows = worksheet.iter_rows(values_only=True)
        headers: list[str] = []
        for index, row in enumerate(rows, start=1):
            if index < header_row:
                continue
            if index == header_row:
                headers = [str(cell).strip() if cell is not None else "" for cell in row]
                continue
            if not any(cell is not None and str(cell).strip() for cell in row):
                continue
            yield {
                headers[i]: row[i]
                for i in range(min(len(headers), len(row)))
                if headers[i]
            }

    def _parse_csv(self, payload: bytes) -> Iterator[dict[str, Any]]:
        text = payload.decode(self.config.get("encoding", "utf-8-sig"))
        delimiter = self.config.get("delimiter")
        if not delimiter:
            delimiter = self._detect_csv_delimiter(text)
        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        for row in reader:
            yield {str(k).strip(): v for k, v in row.items() if k}

    @staticmethod
    def _detect_csv_delimiter(text: str) -> str:
        sample = text[:4096]
        if not sample:
            return ","
        possible = [",", ";", "\t", "|"]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters="".join(possible))
            return dialect.delimiter
        except csv.Error:
            return ";" if ";" in sample and "," not in sample else ","

from __future__ import annotations

import io
from collections.abc import Iterator
from typing import Any

from openpyxl import load_workbook

from apps.ingestion.dto import NormalizedItem
from apps.suppliers.adapters.base import AdapterError, BaseAdapter
from apps.suppliers.adapters.korda import read_product_info
from apps.suppliers.adapters.registry import register_adapter


@register_adapter
class KordaInfoAdapter(BaseAdapter):
    """Korda produktų-info failas (lapai = kalbos kodai) – vertimai.

    Naudojamas kalbiniams feed'ams (feed.language != bazinė) – pipeline
    rašo tik OfferText įrašus. Bazinė kalba (EN) sujungiama tiesiai
    KordaAdapter per config 'info_file'.
    """

    code = "korda_info"

    def parse(self, payload: bytes) -> Iterator[dict[str, Any]]:
        sheet_name = ""
        if getattr(self.feed, "language", None) and self.feed.language:
            sheet_name = (self.feed.language.code or "").upper()
        # config 'files' – keli info workbook'ai (carpporter/guru/korda.xlsx);
        # kitaip – vienas feed.file payload'as.
        files = (getattr(self, "config", None) or {}).get("files")
        if files:
            from pathlib import Path

            from django.conf import settings

            for rel in files:
                path = Path(str(rel))
                if not path.is_absolute():
                    path = Path(settings.BASE_DIR) / path
                if not path.exists():
                    continue
                yield from self._rows(path, sheet_name)
        else:
            yield from self._rows(io.BytesIO(payload), sheet_name)

    @staticmethod
    def _rows(source: Any, sheet_name: str) -> Iterator[dict[str, Any]]:
        workbook = load_workbook(source, read_only=True, data_only=True)
        sheet = sheet_name if sheet_name in workbook.sheetnames else workbook.sheetnames[0]
        if hasattr(source, "seek"):
            source.seek(0)
        for ref, row in read_product_info(source, sheet).items():
            row["ref"] = ref
            yield row

    def normalize(self, row: dict[str, Any]) -> NormalizedItem:
        ref = self._text(row.get("ref"))
        if not ref:
            raise AdapterError(f"Korda info: trūksta Product Ref: {repr(row)[:200]}")
        return NormalizedItem(
            external_id=ref,
            name=self._text(row.get("name")),
            description=self._text(row.get("description")),
            images=list(row.get("images") or []),
            extra={"headgroup": row.get("headgroup"), "subgroup": row.get("subgroup")},
        )

    @staticmethod
    def _text(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()

from __future__ import annotations

import io
from collections.abc import Iterator
from datetime import date, datetime
from decimal import Decimal
from typing import Any

import re
from openpyxl import load_workbook

from apps.catalog.attributes import classify_label
from apps.ingestion.dto import NormalizedItem
from apps.suppliers.adapters.base import AdapterError, BaseAdapter
from apps.suppliers.adapters.registry import register_adapter


# FOX/CCM kalbiniai failai turi IŠVERSTAS stulpelių antraštes.
# Žemėlapis: lokalizuota antraštė -> kanoninė (EN) antraštė.
# Papildomas/kitoks mapping'as gali būti per feed.config["header_map"].
FOX_HEADER_ALIASES: dict[str, dict[str, str]] = {
    "pl": {
        "Nr art.": "Art-No.",
        "Grupa produktów": "Product Group",
        "Marka": "Brand",
        "Nazwa": "Name",
        "Cena hurtowa EUR": "Trade EUR",
        "Cena detaliczna EUR": "RRP EUR",
        "Kod kreskowy EAN sztuka": "Barcode Single EAN",
        "Kod kreskowy EAN opakowanie": "Barcode Tradepack EAN",
        "Kategoria": "Category",
        "Podkategoria": "Sub Category",
        "Rodzina": "Family",
        "Data dostępności": "Availability Date",
        "Opis produktu": "Product Description",
        "Opis techniczny": "Techn. Description",
        "Pełny tekst łączony": "Combined Full Text",
        "Rozmiar": "Size",
        "Wysokość jednostki": "Unit Height",
        "Wysokość detaliczna cm": "Retail Height cm",
        "Szerokość detaliczna cm": "Retail Width cm",
        "Głębokość detaliczna cm": "Retail Depth cm",
        "Waga detaliczna g": "Retail Weight gram",
    },
    "de": {
        "Art.-Nr.": "Art-No.",
        "Produktgruppe": "Product Group",
        "Marke": "Brand",
        # 'Name' lieka 'Name'
        "EK-Preis EUR": "Trade EUR",
        "UVP EUR": "RRP EUR",
        "EAN Einzelbarcode": "Barcode Single EAN",
        "EAN Umverpackung": "Barcode Tradepack EAN",
        "Kategorie": "Category",
        "Unterkategorie": "Sub Category",
        "Familie": "Family",
        "Verfügbarkeitsdatum": "Availability Date",
        "Produktbeschreibung": "Product Description",
        "Technische Beschreibung": "Techn. Description",
        "Kombinierter Volltext": "Combined Full Text",
        "Größe": "Size",
        "Einheitshöhe": "Unit Height",
        "Handelshöhe cm": "Retail Height cm",
        "Handelsbreite cm": "Retail Width cm",
        "Handelstiefe cm": "Retail Depth cm",
        "Handelsgewicht g": "Retail Weight gram",
    },
}

# Video stulpeliai lokalizuoti, bet (XX) kalbos kodas lieka: pvz.
# PL: "Film produktowy (EN)", DE: "Produktvideo (DE)". Alias -> kanoninis prefix.
FOX_VIDEO_HEADER_ALIASES: dict[str, dict[str, str]] = {
    "pl": {
        "Film produktowy": "Product Video",
        "Film instruktażowy": "Instructional Video",
        "Film tutorial": "Tutorial Video",
    },
    "de": {
        "Produktvideo": "Product Video",
        "Anleitungsvideo": "Instructional Video",
        "Tutorialvideo": "Tutorial Video",
    },
}


@register_adapter
class FoxAdapter(BaseAdapter):
    code = "fox"

    def _header_map(self) -> dict[str, str]:
        custom = self.config.get("header_map") or {}
        language_code = (
            self.feed.language.code if self.feed.language_id else ""
        )
        return {**FOX_HEADER_ALIASES.get(language_code, {}), **custom}

    def _canonical_header(self, header: str) -> str:
        header_map = self._header_map()
        if header in header_map:
            return header_map[header]
        # Video antraštės turi kalbos sufiksą: "... (XX)"
        video_match = re.match(r"^(.*)\s\(([A-Z]{2})\)$", header or "")
        if video_match:
            prefixes = FOX_VIDEO_HEADER_ALIASES.get(
                self.feed.language.code if self.feed.language_id else "", {}
            )
            canonical = prefixes.get(video_match.group(1).strip())
            if canonical:
                return f"{canonical} ({video_match.group(2)})"
        return header

    def parse(self, payload: bytes) -> Iterator[dict[str, Any]]:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)

        base_rows = self._sheet_rows(workbook, "BASE")
        text_rows = self._sheet_index(workbook, "TEXT")
        media_cloud_rows = self._sheet_index(workbook, "MEDIA CLOUD LINKS")
        specs_rows = self._sheet_index(workbook, "SPECS")
        video_rows = self._sheet_index(workbook, "VIDEOS")
        compliance_rows = self._sheet_index(workbook, "COMPLIANCE")

        group_brand_map: dict[str, str] = {}
        group_base_name_map: dict[str, str] = {}
        for row in base_rows:
            group_code = self._text(row.get("Product Group"))
            brand = self._text(row.get("Brand"))
            if group_code and brand and group_code not in group_brand_map:
                group_brand_map[group_code] = brand
        for row in text_rows.values():
            group_code = self._text(row.get("Product Group"))
            base_name = self._text(row.get("Product Description"))
            if group_code and base_name and group_code not in group_base_name_map:
                group_base_name_map[group_code] = base_name

        for base in base_rows:
            art_no = self._text(base.get("Art-No."))
            if not art_no:
                continue

            product_group = self._text(base.get("Product Group"))
            text_row = text_rows.get(art_no, {})
            specs_row = specs_rows.get(art_no, {})
            media_cloud_row = media_cloud_rows.get(art_no, {})
            compliance_row = compliance_rows.get(art_no, {})
            videos = self._extract_videos(video_rows.get(art_no, {}))

            variant_name = self._text(base.get("Name"))
            base_name = self._text(text_row.get("Product Description")) or group_base_name_map.get(
                product_group,
                "",
            )
            brand_name = (
                self._text(base.get("Brand"))
                or self._text(text_row.get("Brand"))
                or group_brand_map.get(product_group, "")
            )
            # Pavadinimui visada priešdedamas brand (jei jo dar nėra).
            name = variant_name or base_name or art_no
            if brand_name:
                if not name.lower().startswith(brand_name.lower()):
                    name = f"{brand_name} {name}"
                if base_name and not base_name.lower().startswith(brand_name.lower()):
                    base_name = f"{brand_name} {base_name}"

            category_name = self._text(base.get("Category"))
            sub_category = self._text(base.get("Sub Category"))
            family = self._text(base.get("Family"))

            images = self._extract_media_links(media_cloud_row)
            main_image_url = images[0] if images else ""

            technical_description = self._text(text_row.get("Techn. Description"))
            combined_full_text = self._text(text_row.get("Combined Full Text"))

            axes, features = self._specs_classified(specs_row)
            compliance = self._compliance_attributes(compliance_row)
            if compliance.get("material_fabric"):
                features.setdefault("material_fabric", compliance["material_fabric"])
            size_raw = self._text(specs_row.get("Size"))
            shipping_height_cm = self._to_decimal(specs_row.get("Retail Height cm"))
            shipping_width_cm = self._to_decimal(specs_row.get("Retail Width cm"))
            shipping_depth_cm = self._to_decimal(specs_row.get("Retail Depth cm"))
            shipping_weight_kg = self._grams_to_kg(specs_row.get("Retail Weight gram"))

            identifiers = {
                "art_no": art_no,
                "product_group": product_group,
                "barcode_single_ean": self._text(base.get("Barcode Single EAN")),
                "barcode_tradepack_ean": self._text(base.get("Barcode Tradepack EAN")),
            }

            raw_joined = {
                "base": self._clean_dict(base),
                "text": self._clean_dict(text_row),
                "specs": self._clean_dict(specs_row),
                "media_cloud_links": self._clean_dict(media_cloud_row),
                "videos": self._clean_dict(video_rows.get(art_no, {})),
            }

            yield {
                "external_id": art_no,
                "sku": art_no,
                "name": name,
                "base_name": base_name,
                "description": combined_full_text or base_name or variant_name,
                "technical_description": technical_description,
                "combined_full_text": combined_full_text,
                "brand_name": brand_name,
                "category_path": category_name,
                "sub_category_raw": sub_category,
                "family_raw": family,
                "ean": identifiers["barcode_single_ean"],
                "purchase_price": self._to_decimal(base.get("Trade EUR")),
                "rrp_price": self._to_decimal(base.get("RRP EUR")),
                "currency": "EUR",
                "stock": 0,
                "images": images,
                "main_image_url": main_image_url,
                "videos": videos,
                "attributes": {
                    "axes": axes,
                    "features": features,
                    "compliance": compliance,
                    # be product_group -> atsarginės dalys, jų nereikia
                    "spare_part": not bool(product_group),
                },
                "identifiers": identifiers,
                "size_raw": size_raw,
                "shipping_height_cm": shipping_height_cm,
                "shipping_width_cm": shipping_width_cm,
                "shipping_depth_cm": shipping_depth_cm,
                "shipping_weight_kg": shipping_weight_kg,
                "raw": raw_joined,
            }

    def normalize(self, row: dict[str, Any]) -> NormalizedItem:
        external_id = self._text(row.get("external_id"))
        name = self._text(row.get("name"))
        if not external_id:
            raise AdapterError(f"FOX: trūksta Art-No. eilutėje: {repr(row)[:200]}")
        if not name:
            raise AdapterError(f"FOX: trūksta Name (Art-No.={external_id})")

        attributes = row.get("attributes") if isinstance(row.get("attributes"), dict) else {}
        extra = {
            "identifiers": row.get("identifiers", {}),
            "base_name": self._text(row.get("base_name")),
            "size_raw": self._text(row.get("size_raw")),
            "shipping_height_cm": self._text(row.get("shipping_height_cm")),
            "shipping_width_cm": self._text(row.get("shipping_width_cm")),
            "shipping_depth_cm": self._text(row.get("shipping_depth_cm")),
            "shipping_weight_kg": self._text(row.get("shipping_weight_kg")),
            "product_group": self._text(row.get("identifiers", {}).get("product_group"))
            if isinstance(row.get("identifiers"), dict)
            else "",
            "sub_category_raw": self._text(row.get("sub_category_raw")),
            "family_raw": self._text(row.get("family_raw")),
            "main_image_url": self._text(row.get("main_image_url")),
            "technical_description": self._text(row.get("technical_description")),
            "combined_full_text": self._text(row.get("combined_full_text")),
            "rrp_price": self._to_decimal(row.get("rrp_price")),
        }

        filtered_attributes = self._filter_attributes(attributes)

        return NormalizedItem(
            external_id=external_id,
            sku=self._text(row.get("sku")) or external_id,
            name=name,
            ean=self._text(row.get("ean")),
            brand_name=self._text(row.get("brand_name")),
            category_path=self._text(row.get("category_path")),
            description=self._text(row.get("description")),
            attributes=filtered_attributes,
            images=self._clean_images(row.get("images")),
            videos=self._clean_videos(row.get("videos")),
            stock=0,
            purchase_price=self._to_decimal(row.get("purchase_price")),
            currency=self._text(row.get("currency")) or "EUR",
            is_available=True,
            extra=self._clean_dict(extra),
            raw=self._clean_dict(row.get("raw", {})),
        )

    def _sheet_rows(self, workbook, sheet_name: str) -> list[dict[str, Any]]:
        if sheet_name not in workbook.sheetnames:
            return []
        ws = workbook[sheet_name]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if header is None:
            return []
        headers = [self._canonical_header(self._text(value)) for value in header]
        result: list[dict[str, Any]] = []
        for row in rows:
            if not any(cell not in (None, "") for cell in row):
                continue
            item: dict[str, Any] = {}
            for index, value in enumerate(row):
                if index >= len(headers):
                    break
                key = headers[index]
                if not key:
                    continue
                item[key] = value
            if item:
                result.append(item)
        return result

    def _sheet_index(self, workbook, sheet_name: str) -> dict[str, dict[str, Any]]:
        indexed: dict[str, dict[str, Any]] = {}
        for row in self._sheet_rows(workbook, sheet_name):
            art_no = self._text(row.get("Art-No."))
            if art_no:
                indexed[art_no] = row
        return indexed

    def _extract_media_links(self, media_row: dict[str, Any]) -> list[str]:
        urls: list[str] = []
        for key, value in media_row.items():
            if key == "Art-No.":
                continue
            text = self._text(value)
            if text.lower().startswith("http://") or text.lower().startswith("https://"):
                urls.append(text)
        seen: set[str] = set()
        ordered: list[str] = []
        for url in urls:
            if url in seen:
                continue
            seen.add(url)
            ordered.append(url)
        return ordered

    # "Product Video (EN)" / "Instructional Video (DE)" / "Tutorial Video (PL)"
    _VIDEO_COLUMN_RE = re.compile(r"^(Product|Instructional|Tutorial) Video \(([A-Z]{2})\)$")

    def _extract_videos(self, video_row: dict[str, Any]) -> list[dict[str, str]]:
        videos: list[dict[str, str]] = []
        seen: set[str] = set()
        for key, value in video_row.items():
            match = self._VIDEO_COLUMN_RE.match(str(key).strip())
            if not match:
                continue
            url = self._text(value)
            if not url or not url.lower().startswith(("http://", "https://")) or url in seen:
                continue
            seen.add(url)
            videos.append(
                {
                    "url": url,
                    "type": match.group(1).lower(),
                    "language": match.group(2).lower(),
                }
            )
        return videos

    def _clean_videos(self, value: Any) -> list[dict[str, str]]:
        if not isinstance(value, list):
            return []
        cleaned: list[dict[str, str]] = []
        for item in value:
            if not isinstance(item, dict):
                continue
            url = self._text(item.get("url"))
            if not url:
                continue
            cleaned.append(
                {
                    "url": url,
                    "type": self._text(item.get("type")) or "product",
                    "language": self._text(item.get("language")),
                }
            )
        return cleaned

    def _specs_classified(self, specs_row: dict[str, Any]) -> tuple[dict, dict]:
        """SPEC lapa -> ({axes: code->value}, {features: code->value}).

        Klasifikuoja per kanoninį registrą (visų kalbų label'iai),
        shipping ir ignore stulpeliai praleidžiami.
        """
        axes: dict[str, Any] = {}
        features: dict[str, Any] = {}
        for label, value in (specs_row or {}).items():
            clean = self._clean_value(value)
            if clean in (None, ""):
                continue
            code, kind = classify_label(str(label))
            if not code:
                continue
            if kind == "axis":
                axes.setdefault(code, clean)
            elif kind == "feature":
                features.setdefault(code, clean)
        return axes, features

    def _compliance_attributes(self, compliance_row: dict[str, Any]) -> dict[str, Any]:
        """COMPLIANCE lapa -> material_fabric + papildoma info (gpsr, ingredients)."""
        result: dict[str, Any] = {}
        for label, value in (compliance_row or {}).items():
            clean = self._clean_value(value)
            if clean in (None, ""):
                continue
            low = str(label).strip().lower()
            if "material" in low or "fabric" in low or "materiał" in low or "materialzusammensetzung" in low:
                result.setdefault("material_fabric", clean)
            elif "gpsr" in low or "safety regulation" in low:
                result.setdefault("gpsr", clean)
            elif "ingredient" in low or "składnik" in low or "inhaltsstoff" in low:
                result.setdefault("ingredients", clean)
        return result

    def _grams_to_kg(self, value: Any) -> Decimal | None:
        grams = self._to_decimal(value)
        if grams is None:
            return None
        return grams / Decimal("1000")

    @staticmethod
    def _filter_attributes(attributes: dict[str, Any]) -> dict[str, Any]:
        cleaned = {} if not attributes else {str(key): value for key, value in attributes.items()}
        allowed = {"axes", "features", "compliance", "spare_part", "specs"}
        for key in list(cleaned):
            if key not in allowed:
                cleaned.pop(key, None)
        return cleaned

    def _clean_images(self, value: Any) -> list[str]:
        if isinstance(value, list):
            return [self._text(item) for item in value if self._text(item)]
        return []

    def _clean_dict(self, value: Any) -> dict[str, Any]:
        if not isinstance(value, dict):
            return {}
        cleaned: dict[str, Any] = {}
        for key, item in value.items():
            cleaned[str(key)] = self._clean_value(item)
        return cleaned

    def _clean_value(self, value: Any) -> Any:
        if value is None:
            return None
        if isinstance(value, (bool, int, float, str)):
            return value
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, (date, datetime)):
            return value.isoformat()
        if isinstance(value, list):
            return [self._clean_value(item) for item in value]
        if isinstance(value, tuple):
            return [self._clean_value(item) for item in value]
        if isinstance(value, dict):
            return {str(k): self._clean_value(v) for k, v in value.items()}
        return str(value)

    @staticmethod
    def _text(value: Any) -> str:
        if value is None:
            return ""
        return str(value).strip()
